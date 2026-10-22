"""Спільна бібліотека формування звіту (PDF + DOCX) за вимогами методички."""
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, PageBreak, Paragraph, Preformatted,
                                SimpleDocTemplate, Spacer, Table, TableStyle)
from reportlab.lib import colors
from xml.sax.saxutils import escape

FONT_DIR = "/System/Library/Fonts/Supplemental/"
# шрифти з підтримкою кирилиці
pdfmetrics.registerFont(TTFont("TNR", FONT_DIR + "Times New Roman.ttf"))
pdfmetrics.registerFont(TTFont("TNR-B", FONT_DIR + "Times New Roman Bold.ttf"))
pdfmetrics.registerFont(TTFont("TNR-I", FONT_DIR + "Times New Roman Italic.ttf"))
pdfmetrics.registerFont(TTFont("Mono", FONT_DIR + "Courier New.ttf"))
pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-B", italic="TNR-I", boldItalic="TNR-B")

MAX_LINES = 28
MAX_COLS = 95


def read_evidence(path, max_lines=MAX_LINES):
    """Читає файл виводу та скорочує надто довгі рядки/файли."""
    lines = Path(path).read_text(encoding="utf-8", errors="replace").rstrip("\n").split("\n")
    out = []
    for ln in lines:
        while len(ln) > MAX_COLS:
            out.append(ln[:MAX_COLS])
            ln = "  " + ln[MAX_COLS:]
        out.append(ln)
    if len(out) > max_lines:
        out = out[:max_lines] + [f"... (скорочено, усього {len(out)} рядків)"]
    return "\n".join(out)


class Report:
    def __init__(self, number, topic, author, group, teacher, outdir):
        self.number, self.topic = number, topic
        self.author, self.group, self.teacher = author, group, teacher
        self.outdir = Path(outdir)
        self.items = []

    # --- наповнення ---
    def h1(self, t): self.items.append(("h1", t))
    def h2(self, t): self.items.append(("h2", t))
    def p(self, t): self.items.append(("p", t))
    def bullets(self, lst): self.items.append(("bullets", lst))
    def code(self, t, cap=None): self.items.append(("code", t, cap))
    def table(self, header, rows, widths=None): self.items.append(("table", header, rows, widths))
    def img(self, path, cap, width_cm=15): self.items.append(("img", path, cap, width_cm))
    def pagebreak(self): self.items.append(("pb",))

    def step(self, title, cmd_text, evidence, interp, max_lines=MAX_LINES):
        """Блок «команда → результат → інтерпретація»."""
        self.items.append(("step", title, cmd_text, evidence, interp, max_lines))

    # --- PDF ---
    def build_pdf(self, path):
        base = ParagraphStyle("b", fontName="TNR", fontSize=12, leading=14.4,
                              firstLineIndent=1.25 * cm, alignment=TA_JUSTIFY)
        noind = ParagraphStyle("n", parent=base, firstLineIndent=0)
        center = ParagraphStyle("c", parent=base, alignment=TA_CENTER, firstLineIndent=0)
        centerb = ParagraphStyle("cb", parent=center, fontName="TNR-B")
        h1 = ParagraphStyle("h1", parent=noind, fontName="TNR-B", spaceBefore=10, spaceAfter=6)
        h2 = ParagraphStyle("h2", parent=noind, fontName="TNR-B", spaceBefore=6, spaceAfter=4)
        cap = ParagraphStyle("cap", parent=center, fontName="TNR-I", fontSize=11, leading=13)
        mono = ParagraphStyle("m", fontName="Mono", fontSize=8, leading=9.4,
                              backColor=colors.Color(0.95, 0.95, 0.95), leftIndent=2, rightIndent=2)
        cell = ParagraphStyle("cell", parent=noind, fontSize=10.5, leading=12.5, alignment=0)

        def para(t, st=base): return Paragraph(escape(t).replace("\n", "<br/>"), st)

        fl = []
        fl += [Spacer(1, 1 * cm), para("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", centerb),
               para("Національний технічний університет України", center),
               para("«Київський політехнічний інститут імені Ігоря Сікорського»", center),
               Spacer(1, 5 * cm), para(f"Лабораторна робота № {self.number}", centerb),
               para("з дисципліни «Сучасні інженерні мережеві технології»", center),
               Spacer(1, 0.5 * cm), para(f"Тема: {self.topic}", center),
               Spacer(1, 6 * cm)]
        sign = Table([[para("Виконав(ла):", noind), para(f"{self.author}, група {self.group}", noind)],
                      [para("Роботу прийняв:", noind), para(self.teacher, noind)]],
                     colWidths=[4 * cm, 11 * cm])
        fl += [sign, Spacer(1, 3 * cm), para("Київ — 2026", center), PageBreak()]

        for it in self.items:
            k = it[0]
            if k == "h1": fl.append(para(it[1], h1))
            elif k == "h2": fl.append(para(it[1], h2))
            elif k == "p": fl.append(para(it[1], base))
            elif k == "bullets":
                for b in it[1]: fl.append(para("• " + b, ParagraphStyle("bl", parent=noind, leftIndent=1.25 * cm, firstLineIndent=-0.5 * cm)))
            elif k == "code":
                fl.append(Preformatted(it[1], mono))
                if it[2]: fl.append(para(it[2], cap))
                fl.append(Spacer(1, 4))
            elif k == "table":
                _, header, rows, widths = it
                data = [[para(str(c), ParagraphStyle("th", parent=cell, fontName="TNR-B")) for c in header]]
                data += [[para(str(c), cell) for c in r] for r in rows]
                t = Table(data, colWidths=[w * cm for w in widths] if widths else None, repeatRows=1)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.9, 0.9, 0.9)),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                fl += [t, Spacer(1, 6)]
            elif k == "img":
                from PIL import Image as PI
                w, h = PI.open(it[1]).size
                wcm = it[3]
                fl += [Image(str(it[1]), width=wcm * cm, height=wcm * cm * h / w), para(it[2], cap), Spacer(1, 4)]
            elif k == "pb": fl.append(PageBreak())
            elif k == "step":
                _, title, cmd, ev, interp, ml = it
                fl.append(para(title, h2))
                fl.append(para("Команда:", noind))
                fl.append(Preformatted("\n".join(_wrap(cmd)), mono))
                fl.append(para("Результат:", noind))
                fl.append(Preformatted(read_evidence(ev, ml) if ev else "—", mono))
                fl.append(para("Інтерпретація: " + interp, noind))
                fl.append(Spacer(1, 6))

        def footer(c, d):
            c.setFont("TNR", 10)
            c.drawCentredString(A4[0] / 2, 1 * cm, str(d.page))

        doc = SimpleDocTemplate(str(path), pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                                leftMargin=2.5 * cm, rightMargin=2 * cm,
                                title=f"ЛР{self.number}", author=self.author)
        doc.build(fl, onLaterPages=footer)

    # --- DOCX ---
    def build_docx(self, path):
        d = Document()
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin = s.bottom_margin = s.right_margin = Cm(2)
        s.left_margin = Cm(2.5)
        st = d.styles["Normal"]
        st.font.name, st.font.size = "Times New Roman", Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.space_after = Pt(0)

        def P(t, bold=False, center=False, indent=True, italic=False):
            p = d.add_paragraph()
            r = p.add_run(t); r.bold, r.italic = bold, italic
            if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif indent: p.paragraph_format.first_line_indent = Cm(1.25)
            return p

        def MONO(t):
            p = d.add_paragraph()
            r = p.add_run(t); r.font.name = "Courier New"; r.font.size = Pt(8)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            return p

        for _ in range(2): P("")
        P("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", True, True)
        P("Національний технічний університет України", center=True)
        P("«Київський політехнічний інститут імені Ігоря Сікорського»", center=True)
        for _ in range(8): P("")
        P(f"Лабораторна робота № {self.number}", True, True)
        P("з дисципліни «Сучасні інженерні мережеві технології»", center=True)
        P(f"Тема: {self.topic}", center=True)
        for _ in range(9): P("")
        P(f"Виконав(ла): {self.author}, група {self.group}", indent=False)
        P(f"Роботу прийняв: {self.teacher}", indent=False)
        for _ in range(4): P("")
        P("Київ — 2026", center=True)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        for it in self.items:
            k = it[0]
            if k in ("h1", "h2"): P(it[1], True, indent=False)
            elif k == "p": P(it[1])
            elif k == "bullets":
                for b in it[1]: P("• " + b, indent=False)
            elif k == "code":
                MONO(it[1])
                if it[2]: P(it[2], center=True, italic=True)
            elif k == "table":
                _, header, rows, widths = it
                t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
                for i, h in enumerate(header):
                    t.rows[0].cells[i].text = ""
                    t.rows[0].cells[i].paragraphs[0].add_run(str(h)).bold = True
                for r in rows:
                    cells = t.add_row().cells
                    for i, c in enumerate(r): cells[i].text = str(c)
                for row in t.rows:
                    for c in row.cells:
                        for pp in c.paragraphs:
                            for rr in pp.runs: rr.font.size = Pt(10.5)
                d.add_paragraph()
            elif k == "img":
                d.add_picture(str(it[1]), width=Cm(min(it[3], 16)))
                d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                P(it[2], center=True, italic=True)
            elif k == "pb": d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif k == "step":
                _, title, cmd, ev, interp, ml = it
                P(title, True, indent=False)
                P("Команда:", indent=False); MONO("\n".join(_wrap(cmd)))
                P("Результат:", indent=False); MONO(read_evidence(ev, ml) if ev else "—")
                P("Інтерпретація: " + interp, indent=False)
        d.save(str(path))

    def build(self):
        self.outdir.mkdir(parents=True, exist_ok=True)
        stem = f"{self.author.replace(' ', '_')}_{self.group}_ЛР{self.number}"
        self.build_pdf(self.outdir / f"{stem}.pdf")
        self.build_docx(self.outdir / f"{stem}.docx")
        return stem


def _wrap(text, width=MAX_COLS):
    out = []
    for ln in text.split("\n"):
        while len(ln) > width:
            out.append(ln[:width]); ln = "  " + ln[width:]
        out.append(ln)
    return out
